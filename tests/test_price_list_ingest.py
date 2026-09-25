"""Тесты разбора прайс-листа питомника (pipeline/catalog/price_list_ingest.py)
— прямой запрос пользователя (2026-09-19): «нужна загрузка для прайс-листа с
распознаванием растений = коррекция списка, вшитого в сервис»."""

from __future__ import annotations

import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parents[1]))

import openpyxl
import pytest
from docx import Document

from pipeline.catalog.dpioos_species_catalog import DpioosSpeciesCatalog
from pipeline.catalog.price_list_ingest import (
    ParsedPriceRow,
    UnsupportedPriceListFormatError,
    build_proposal,
    classify_row,
    parse_price_list,
)


@pytest.fixture(scope="module")
def catalog():
    return DpioosSpeciesCatalog()


def test_parse_excel_extracts_name_and_price(tmp_path):
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.append(["Ель обыкновенная, крупномер", 25000])
    ws.append(["Липа мелколистная", 1800])
    ws.append(["Пусто", None])
    path = tmp_path / "price.xlsx"
    wb.save(path)

    rows = parse_price_list(path)
    assert len(rows) == 2
    assert rows[0].name_text == "Ель обыкновенная, крупномер"
    assert rows[0].price_rub == 25000
    assert rows[1].price_rub == 1800


def test_parse_docx_table_extracts_name_and_price(tmp_path):
    doc = Document()
    table = doc.add_table(rows=2, cols=2)
    table.rows[0].cells[0].text = "Можжевельник казацкий"
    table.rows[0].cells[1].text = "3500 руб"
    table.rows[1].cells[0].text = "Кустарник спирея"
    table.rows[1].cells[1].text = "800"
    path = tmp_path / "price.docx"
    doc.save(path)

    rows = parse_price_list(path)
    assert len(rows) == 2
    names = [r.name_text for r in rows]
    assert "Можжевельник казацкий" in names
    prices = {r.name_text: r.price_rub for r in rows}
    assert prices["Можжевельник казацкий"] == 3500


def test_parse_docx_paragraph_fallback_without_table(tmp_path):
    doc = Document()
    doc.add_paragraph("Сосна обыкновенная - 4500 руб")
    doc.add_paragraph("Просто текст без цены")
    path = tmp_path / "price_no_table.docx"
    doc.save(path)

    rows = parse_price_list(path)
    assert len(rows) == 1
    assert rows[0].price_rub == 4500


def test_parse_unsupported_extension_raises(tmp_path):
    bad = tmp_path / "price.txt"
    bad.write_text("Ель - 1000 руб", encoding="utf-8")
    with pytest.raises(UnsupportedPriceListFormatError):
        parse_price_list(bad)


def test_classify_row_exact_catalog_match(catalog):
    entry = catalog.all_entries()[0]
    row = ParsedPriceRow(name_text=entry.name_ru, price_rub=1000, source_location="test")
    result = classify_row(row, catalog)
    assert result.matched_species_name == entry.name_ru
    assert result.category_key is not None


def test_classify_row_keyword_fallback_conifer_tree(catalog):
    row = ParsedPriceRow(name_text="Ель голубая (не из каталога, вымышленный сорт)", price_rub=20000, source_location="test")
    result = classify_row(row, catalog)
    assert result.matched_species_name is None
    assert result.category_key == "conifer_tree_standard"


def test_classify_row_keyword_fallback_shrub_large(catalog):
    row = ParsedPriceRow(name_text="Кустарник дерен белый, крупномер", price_rub=2500, source_location="test")
    result = classify_row(row, catalog)
    assert result.category_key == "deciduous_shrub_large"


def test_classify_row_unclassified_when_no_keywords_match(catalog):
    row = ParsedPriceRow(name_text="Загадочное растение XYZ", price_rub=999, source_location="test")
    result = classify_row(row, catalog)
    assert result.category_key is None


def test_classify_row_ordinary_deciduous_tree_by_genus(catalog):
    # Реальный баг, найденный на живой проверке (2026-09-20): раньше
    # классификация требовала буквальное слово "дерев" в тексте — из-за чего
    # обычные названия видов (без слова "дерево") неверно попадали в
    # unclassified. Родовое слово "липа" должно определяться напрямую из
    # каталога ДПиООС, без точного совпадения полного названия.
    row = ParsedPriceRow(name_text="Липа мелколистная 30-40", price_rub=1900, source_location="test")
    result = classify_row(row, catalog)
    assert result.category_key == "deciduous_tree_standard"


def test_classify_row_ordinary_deciduous_tree_large_by_genus(catalog):
    row = ParsedPriceRow(name_text="Береза повислая, крупномер", price_rub=12000, source_location="test")
    result = classify_row(row, catalog)
    assert result.category_key == "deciduous_tree_large"


def test_build_proposal_aggregates_median_for_existing_category():
    from pipeline.catalog.price_list_ingest import ClassifiedRow

    rows = [
        ClassifiedRow(ParsedPriceRow("Ель А", 10000, "t"), "conifer_tree_standard", None, False),
        ClassifiedRow(ParsedPriceRow("Ель Б", 20000, "t"), "conifer_tree_standard", None, False),
        ClassifiedRow(ParsedPriceRow("Ель В", 90000, "t"), "conifer_tree_standard", None, False),
    ]
    current = {"conifer_tree_standard": {"default_rub": 6000, "min_rub": 2000, "max_rub": 15000}}
    proposals = build_proposal(rows, current)
    assert len(proposals) == 1
    p = proposals[0]
    assert p.is_new is False
    assert p.current_default_rub == 6000
    assert p.proposed_default_rub == 20000  # медиана 10000/20000/90000
    assert p.matched_row_count == 3


def test_build_proposal_marks_new_category():
    from pipeline.catalog.price_list_ingest import ClassifiedRow

    rows = [ClassifiedRow(ParsedPriceRow("Что-то новое", 5000, "t"), "conifer_shrub", None, False)]
    proposals = build_proposal(rows, current_categories={})
    assert proposals[0].is_new is True
    assert proposals[0].current_default_rub is None
